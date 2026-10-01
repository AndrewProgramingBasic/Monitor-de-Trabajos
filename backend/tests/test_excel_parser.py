import io
import pytest
from datetime import datetime
import openpyxl
import pytz
from services.excel_parser import (
    clean_cdc_number,
    check_rescheduled,
    find_header_row,
    map_headers,
    normalize_header,
    normalize_spanish_datetime,
    parse_committee_sheet,
    strip_accents,
)


def test_strip_accents():
    assert strip_accents("Miércoles") == "Miercoles"
    assert strip_accents("Afectación") == "Afectacion"
    assert strip_accents("Título") == "Titulo"


def test_normalize_spanish_datetime_formats():
    tz = pytz.timezone("America/Caracas")

    # Sample 1
    dt1 = normalize_spanish_datetime("Lunes 28/9/2026 15:20")
    assert dt1 is not None
    assert dt1.year == 2026 and dt1.month == 9 and dt1.day == 28
    assert dt1.hour == 15 and dt1.minute == 20
    assert dt1.tzinfo.zone == "America/Caracas"

    # Sample 2
    dt2 = normalize_spanish_datetime("Viernes 2/10/2026 3:10")
    assert dt2 is not None
    assert dt2.year == 2026 and dt2.month == 10 and dt2.day == 2
    assert dt2.hour == 3 and dt2.minute == 10

    # Sample 3
    dt3 = normalize_spanish_datetime("Miércoles 30/09/2026 00:50")
    assert dt3 is not None
    assert dt3.year == 2026 and dt3.month == 9 and dt3.day == 30
    assert dt3.hour == 0 and dt3.minute == 50

    # Sample 4 (Period instead of colon)
    dt4 = normalize_spanish_datetime("Miercoles 30/9/2026 16.10 ")
    assert dt4 is not None
    assert dt4.year == 2026 and dt4.month == 9 and dt4.day == 30
    assert dt4.hour == 16 and dt4.minute == 10

    # Sample 5
    dt5 = normalize_spanish_datetime("Viernes 02/10/2025 00:55")
    assert dt5 is not None
    assert dt5.year == 2025 and dt5.month == 10 and dt5.day == 2
    assert dt5.hour == 0 and dt5.minute == 55

    # Native datetime
    native_dt = datetime(2026, 9, 24, 10, 30)
    dt6 = normalize_spanish_datetime(native_dt)
    assert dt6.tzinfo.zone == "America/Caracas"
    assert dt6.hour == 10 and dt6.minute == 30


def test_normalize_datetime_errors():
    # Optional field returns None
    assert normalize_spanish_datetime(None, is_mandatory=False) is None
    assert normalize_spanish_datetime("", is_mandatory=False) is None

    # Mandatory field raises ValueError
    with pytest.raises(ValueError, match="no puede estar vacío"):
        normalize_spanish_datetime(None, is_mandatory=True, field_name="start_datetime", row_idx=5)

    with pytest.raises(ValueError, match="Formato de fecha y hora inválido"):
        normalize_spanish_datetime("texto_no_valido", is_mandatory=True, field_name="start_datetime", row_idx=5)


def test_clean_cdc_number():
    assert clean_cdc_number(" 3900092507 ") == "3900092507"
    assert clean_cdc_number(3900092507) == "3900092507"
    assert clean_cdc_number("3900092507.0") == "3900092507"
    assert clean_cdc_number("R/P 3900091216") == "R/P 3900091216"
    assert clean_cdc_number("") is None
    assert clean_cdc_number("null") is None


def test_check_rescheduled_logic():
    class MockCursor:
        def __init__(self, existing_rows=None):
            self.existing_rows = existing_rows or {}

        def execute(self, query, params):
            self.last_params = params

        def fetchone(self):
            cdc = self.last_params[0]
            if cdc in self.existing_rows:
                return (self.existing_rows[cdc],)
            return None

    # Case 1: CDC with R/P and matches existing DB row
    cur = MockCursor(existing_rows={"3900091216": 42})
    is_resched, parent_id = check_rescheduled(cur, "R/P 3900091216", {})
    assert is_resched is True
    assert parent_id == 42

    # Case 2: CDC with R/P but no prior DB row yet
    cur = MockCursor(existing_rows={})
    is_resched, parent_id = check_rescheduled(cur, "R/P 3900099999", {})
    assert is_resched is True
    assert parent_id is None

    # Case 3: CDC without R/P, but exists in DB
    cur = MockCursor(existing_rows={"3900092507": 10})
    is_resched, parent_id = check_rescheduled(cur, "3900092507", {})
    assert is_resched is True
    assert parent_id == 10

    # Case 4: Brand new normal CDC
    cur = MockCursor(existing_rows={})
    is_resched, parent_id = check_rescheduled(cur, "3900092507", {})
    assert is_resched is False
    assert parent_id is None

    # Case 5: Seen in current batch
    cur = MockCursor(existing_rows={})
    is_resched, parent_id = check_rescheduled(cur, "3900092507", {"3900092507": 99})
    assert is_resched is True
    assert parent_id == 99


def test_excel_file_parsing_integration(monkeypatch):
    """Generates an in-memory openpyxl workbook matching the specification and parses it."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comite_VPTI"

    # Row 2: Title
    ws.cell(row=2, column=1, value="COMITÉ INTERNO VPTI 24/09/2026")

    # Row 4: Header
    headers = [
        "N°", "Título del Trabajo", "CDC", "Justificación",
        "Fecha y Hora Inicio", "Fecha y Hora Fin", "Afectación",
        "Inicio Afectación", "Fin Afectación", "Detalle Afectación",
        "Aprobación Comité VPTI", "Aprobación Gerencia"
    ]
    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=4, column=col_idx, value=h)

    # Row 5: Task 1
    ws.cell(row=5, column=1, value=1)
    ws.cell(row=5, column=2, value="Actualización de firmware")
    ws.cell(row=5, column=3, value="3900092507")
    ws.cell(row=5, column=4, value="Actualización preventiva")
    ws.cell(row=5, column=5, value="Lunes 28/9/2026 15:20")
    ws.cell(row=5, column=6, value="Lunes 28/9/2026 17:20")
    ws.cell(row=5, column=7, value="NO")
    ws.cell(row=5, column=11, value="Aprobado")
    ws.cell(row=5, column=12, value="Aprobado")

    # Row 6: Task 2 (Rescheduled R/P)
    ws.cell(row=6, column=1, value=2)
    ws.cell(row=6, column=2, value="Migración enlace principal")
    ws.cell(row=6, column=3, value="R/P 3900091216")
    ws.cell(row=6, column=4, value="Migración planificada")
    ws.cell(row=6, column=5, value="Miercoles 30/9/2026 16.10 ")
    ws.cell(row=6, column=6, value="Miercoles 30/9/2026 18:00")
    ws.cell(row=6, column=7, value="SI")
    ws.cell(row=6, column=8, value="Miercoles 30/9/2026 16.10 ")
    ws.cell(row=6, column=9, value="Miercoles 30/9/2026 18:00")
    ws.cell(row=6, column=10, value="Conectividad sede")
    ws.cell(row=6, column=11, value="Aprobado")
    ws.cell(row=6, column=12, value="Pendiente")

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    # Mock database interaction for testing parse logic
    class FakeConnection:
        def __init__(self):
            self.committed = False
            self.rolled_back = False
            self.task_ids = [101, 102]

        def cursor(self):
            return self

        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc_val, exc_tb):
            pass

        def execute(self, query, params=None):
            self.last_query = query
            self.last_params = params

        def fetchone(self):
            if "INSERT INTO committee_sheets" in getattr(self, "last_query", ""):
                return (1,)
            if "SELECT id FROM scheduled_tasks" in getattr(self, "last_query", ""):
                return None  # No prior tasks
            if "INSERT INTO scheduled_tasks" in getattr(self, "last_query", ""):
                return (self.task_ids.pop(0) if self.task_ids else 999,)
            return None

        def commit(self):
            self.committed = True

        def rollback(self):
            self.rolled_back = True

    fake_conn = FakeConnection()
    from contextlib import contextmanager
    @contextmanager
    def mock_get_db_connection():
        yield fake_conn

    monkeypatch.setattr("services.excel_parser.get_db_connection", mock_get_db_connection)

    summary = parse_committee_sheet(stream, "comite_2409.xlsx", uploaded_by_user_id=1)
    assert summary["sheet_id"] == 1
    assert summary["committee_name"] == "COMITÉ INTERNO VPTI 24/09/2026"
    assert summary["tasks_imported"] == 2
    assert summary["rescheduled_count"] == 1  # The R/P task
    assert fake_conn.committed is True


def test_normalize_header():
    assert normalize_header("  Título del\nTrabajo  ") == "titulo del trabajo"
    assert normalize_header("Fecha  inicio del Trabajo ") == "fecha inicio del trabajo"
    assert normalize_header("Afectación\r\n(Inicio)") == "afectacion (inicio)"
    assert normalize_header("Aprobación  de los Gerentes") == "aprobacion de los gerentes"
    assert normalize_header(None) == ""


def test_map_headers_fuzzy():
    # Real production committee headers
    headers = {
        2: "n",
        3: "titulo del trabajo",
        4: "cdc",
        5: "fecha inicio del trabajo",
        6: "fecha fin del trabajo",
        7: "justificacion del trabajo",
        8: "afectacion",
        9: "fecha y hora de afectacion (inicio)",
        10: "fecha y hora de afectacion (fin)",
        11: "detalle de la afectacion",
        12: "aprobado comite interno vpti",
        13: "aprobacion de los gerentes",
    }
    mapping = map_headers(headers)
    assert mapping["title"] == 3
    assert mapping["cdc"] == 4
    assert mapping["start_datetime"] == 5
    assert mapping["end_datetime"] == 6
    assert mapping["justification"] == 7
    assert mapping["has_affectation"] == 8
    assert mapping["affectation_start"] == 9
    assert mapping["affectation_end"] == 10
    assert mapping["affectation_details"] == 11
    assert mapping["vpti_committee_approval"] == 12
    assert mapping["managers_approval"] == 13
    assert mapping["sheet_item_order"] == 2

    # Alternate abbreviations and synonyms
    alt_headers = {
        1: "item",
        2: "descripcion",
        3: "cdc",
        4: "f. inicio",
        5: "f. fin",
        6: "motivo",
        7: "afectacion",
        8: "comite vpti",
        9: "gerente",
    }
    alt_mapping = map_headers(alt_headers)
    assert alt_mapping["sheet_item_order"] == 1
    assert alt_mapping["title"] == 2
    assert alt_mapping["cdc"] == 3
    assert alt_mapping["start_datetime"] == 4
    assert alt_mapping["end_datetime"] == 5
    assert alt_mapping["justification"] == 6
    assert alt_mapping["has_affectation"] == 7
    assert alt_mapping["vpti_committee_approval"] == 8
    assert alt_mapping["managers_approval"] == 9


def test_find_header_row_dynamic():
    # Test workbook where header is placed at row 6
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=2, column=1, value="COMITÉ INTERNO VPTI")
    headers = ["N°", "Título", "CDC", "Fecha Inicio", "Fecha Fin"]
    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=6, column=col_idx, value=h)

    h_row, _, mapping = find_header_row(ws)
    assert h_row == 6
    assert mapping["title"] == 2
    assert mapping["cdc"] == 3
    assert mapping["start_datetime"] == 4
    assert mapping["end_datetime"] == 5


def test_real_file_parser_dry_run():
    from pathlib import Path
    backend_dir = Path(__file__).resolve().parent.parent
    candidates = [
        backend_dir.parent / "CUADRO COMITE 24 09 2026 JUAN GARCIA.XLSX",
        backend_dir / "CUADRO COMITE 24 09 2026 JUAN GARCIA.XLSX",
    ]
    target_file = next((p for p in candidates if p.exists()), None)
    if not target_file:
        pytest.skip("CUADRO COMITE 24 09 2026 JUAN GARCIA.XLSX not found")

    result = parse_committee_sheet(str(target_file), dry_run=True)
    assert result["header_row"] == 4
    assert result["tasks_imported"] == 24
    assert result["rescheduled_count"] >= 1
    assert len(result["tasks"]) == 24
    assert result["tasks"][0]["start_datetime"].tzinfo.zone == "America/Caracas"
